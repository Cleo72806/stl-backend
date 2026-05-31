# ============================================================
# preprocessor.py — Dynamic Column Detection
# Direct Python port of preprocessor.R
# ============================================================

import io
import re
import logging
import os
from datetime import date
from dateutil.relativedelta import relativedelta
from typing import Optional

import pandas as pd
import openpyxl
from openpyxl.utils import column_index_from_string

from storage import save_parquet, rebuild_master

logger = logging.getLogger(__name__)

PALATANDAAN: dict[str, str] = {
    "BCQ (MWH)":                       "bFm",
    "CUSTOMER FM HOURS":                "xFh",
    "CAPACITY PRICE WITHOUT FM (PHP)":  "xFm",
    "CAPACITY PRICE WITH FM (PHP)":     "yFm",
    "ENERGY PRICE (PHP)":               "zFm",
    "CONTRACT PRICE WITH FM(PHP)":      "aFm",
}

SKIP_LABELS = {"TOTAL", "EFFECTIVE RATE", "DATE", "INTERVAL", "SUM", "BCQ"}


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip()).upper()


def _col_letter_to_idx(col_str: str) -> int:
    col_str = re.sub(r"[^A-Za-z]", "", col_str).upper()
    return column_index_from_string(col_str)


def _build_dynamic_map(wb: openpyxl.Workbook, sheet_name: str) -> dict:
    ws = wb[sheet_name]

    row1: dict[str, int] = {}
    row5: dict[int, str] = {}

    for cell in ws[1]:
        if cell.value:
            norm = _normalize(str(cell.value))
            row1[norm] = cell.column

    for cell in ws[5]:
        if cell.value:
            row5[cell.column] = str(cell.value).strip()

    if not row1 or not row5:
        logger.warning(f"Sheet '{sheet_name}': could not read rows 1/5.")
        return {}

    metric_label_cols: dict[str, int] = {}
    for label_norm, code in PALATANDAAN.items():
        if label_norm in row1:
            metric_label_cols[code] = row1[label_norm]

    if not metric_label_cols:
        logger.warning(f"Sheet '{sheet_name}': no palatandaan found in row 1.")
        return {}

    sorted_metrics = sorted(metric_label_cols.items(), key=lambda x: x[1])

    result: dict[str, dict[str, int]] = {}

    for i, (code, label_col) in enumerate(sorted_metrics):
        actual_start = label_col - 4

        total_col = None
        for scan_col in range(actual_start, actual_start + 31):
            name = row5.get(scan_col, "")
            if _normalize(name) == "TOTAL":
                total_col = scan_col
                break

        if total_col is not None:
            actual_end = total_col
        elif i < len(sorted_metrics) - 1:
            next_label_col = sorted_metrics[i + 1][1]
            actual_end = next_label_col - 4 - 1
        else:
            actual_end = max(row5.keys()) if row5 else actual_start + 30

        cust_map: dict[str, int] = {}
        for col_idx, name in row5.items():
            if actual_start <= col_idx <= actual_end:
                name_up = _normalize(name)
                if name_up in SKIP_LABELS:
                    if name_up == "TOTAL":
                        cust_map["All Customers"] = col_idx
                else:
                    cust_map[name] = col_idx

        result[code] = cust_map

    n_cust = len({c for m in result.values() for c in m.keys()})
    logger.info(
        f"Sheet '{sheet_name}': dynamic map OK — "
        f"{len(result)} metric groups, {n_cust} unique customers."
    )
    return result


def _get_col_index(dyn_map: dict, metric_code: str, customer: str) -> Optional[int]:
    metric_map = dyn_map.get(metric_code, {})
    cust_norm = _normalize(customer)
    for name, idx in metric_map.items():
        if _normalize(name) == cust_norm:
            return idx
    logger.warning(f"No column for metric='{metric_code}', customer='{customer}'")
    return None


def _process_col(ws, col_idx: int, start_date: date, data_start_row: int = 7) -> list[dict]:
    if col_idx is None:
        return []

    raw = []
    for row in ws.iter_rows(min_row=data_start_row, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        try:
            raw.append(float(cell.value) if cell.value is not None else 0.0)
        except (TypeError, ValueError):
            raw.append(0.0)

    n_days = len(raw) // 24
    if n_days == 0:
        return []

    records = []
    for d in range(n_days):
        day_val = sum(raw[d * 24: d * 24 + 24])
        day_date = start_date + relativedelta(days=d)
        records.append({"Date": day_date, "Value": day_val})
    return records


def process_uploaded_file(file_bytes: bytes, original_filename: str) -> str:
    fname_upper = original_filename.upper()
    is_gmec = "GMEC" in fname_upper
    is_gnpd = "GNPD" in fname_upper
    if not is_gmec and not is_gnpd:
        is_gmec = is_gnpd = True

    logger.info(f"Processing '{original_filename}'")

    buf = io.BytesIO(file_bytes)
    wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)

    all_records: list[dict] = []

    for sheet_name in wb.sheetnames:
        s = sheet_name.strip()
        if len(s) != 6 or not s.isdigit():
            logger.info(f"Skipping sheet '{sheet_name}' — not a date sheet.")
            continue

        year = int(s[:4])
        month = int(s[4:])
        billing_start = date(year, month, 26) - relativedelta(months=1)

        logger.info(f"Processing sheet '{s}' | billing start: {billing_start}")

        ws = wb[sheet_name]
        dyn_map = _build_dynamic_map(wb, sheet_name)

        all_customers: list[str] = list(
            {c for m in dyn_map.values() for c in m.keys()}
        ) if dyn_map else []

        gmec_metrics = ["xFm", "yFm", "zFm", "aFm", "bFm", "xFh"]
        gnpd_metrics = ["xFm", "yFm"]

        for customer in all_customers:
            for metric_code in (gmec_metrics if is_gmec else []) + (gnpd_metrics if is_gnpd and not is_gmec else []):
                provider = "GMEC" if metric_code in gmec_metrics and is_gmec else "GNPD"
                if not is_gmec and provider == "GMEC":
                    continue
                if not is_gnpd and provider == "GNPD":
                    continue

                col_idx = _get_col_index(dyn_map, metric_code, customer)
                if col_idx is None:
                    continue

                rows = _process_col(ws, col_idx, billing_start)
                for r in rows:
                    all_records.append({
                        "Date":     r["Date"],
                        "Provider": provider,
                        "Customer": customer,
                        "Metric":   metric_code,
                        "Value":    r["Value"],
                    })

    wb.close()

    if not all_records:
        raise ValueError(f"No valid data processed for file: {original_filename}")

    df = pd.DataFrame(all_records)
    df["Date"] = pd.to_datetime(df["Date"])

    output_filename = os.path.splitext(original_filename)[0] + ".parquet"
    save_parquet(df, output_filename)
    logger.info(f"Processed and synced: {output_filename}")

    rebuild_master()

    return output_filename